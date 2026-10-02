import os
import sys
import unittest
import openpyxl
from unittest.mock import patch, MagicMock

# Add backend directory to sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from utils.validation import is_valid_email, normalize_selection
from services.excel_service import parse_excel_file
from services.email_service import substitute_placeholders, send_campaign_emails
from services.sender_service import add_sender, get_all_senders, test_gmail_credentials
from services.campaign_service import (
    create_campaign,
    save_campaign,
    get_campaign_by_id,
    remap_campaign_columns,
    update_campaign_email_content,
    get_email_preview,
    execute_campaign_send
)

class TestEmailAutomationPlatform(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Create test Excel files
        cls.test_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'test_files')
        os.makedirs(cls.test_dir, exist_ok=True)
        
        # Test file 1: Mixed statuses
        cls.file1 = os.path.join(cls.test_dir, 'test_candidates.xlsx')
        wb1 = openpyxl.Workbook()
        ws1 = wb1.active
        ws1.append(['Name', 'Email', 'Selection', 'Department', 'Round'])
        ws1.append(['Ravi Kumar', 'ravi@example.com', 'Selected', 'AI', 'Round 1'])
        ws1.append(['Priya Sharma', 'priya@example.com', 'selected', 'Cloud', 'Round 1'])
        ws1.append(['Vikram Singh', 'vikram@example.com', 'Not Selected', 'Backend', 'Round 1'])
        ws1.append(['Ananya Iyer', 'ananya@example.com', 'NOT SELECTED', 'DevOps', 'Round 1'])
        ws1.append(['Divya Nair', 'divya@example.com', 'Pending', 'QA', 'Round 1']) # Unknown
        ws1.append(['Meera Joshi', '', 'Selected', 'UI', 'Round 1']) # Missing email
        ws1.append(['Karan Invalid', 'not-an-email', 'Selected', 'Sec', 'Round 1']) # Invalid email
        wb1.save(cls.file1)
        
        # Test file 2: Zero selected candidates
        cls.file_zero = os.path.join(cls.test_dir, 'zero_selected.xlsx')
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.append(['Name', 'Email', 'Selection'])
        ws2.append(['Person A', 'a@example.com', 'Not Selected'])
        ws2.append(['Person B', 'b@example.com', 'Pending'])
        wb2.save(cls.file_zero)

    def test_01_selection_and_not_selected_filtering(self):
        """TEST 1: Excel with Selected and Not Selected rows. Only Selected rows are eligible."""
        parsed = parse_excel_file(self.file1)
        summary = parsed['summary']
        self.assertEqual(summary['total_rows'], 7)
        self.assertEqual(summary['selected_count'], 4)
        self.assertEqual(summary['not_selected_count'], 2)
        self.assertEqual(summary['unknown_count'], 1)
        self.assertEqual(summary['missing_email_count'], 1)
        self.assertEqual(summary['valid_selected_recipients'], 2)
        
        # Only Ravi and Priya should be in valid_recipients
        valid_names = [r['name'] for r in parsed['valid_recipients']]
        self.assertIn('Ravi Kumar', valid_names)
        self.assertIn('Priya Sharma', valid_names)
        self.assertNotIn('Vikram Singh', valid_names)
        self.assertNotIn('Ananya Iyer', valid_names)

    def test_02_different_gmail_sender(self):
        """TEST 2: Different Gmail sender accounts configured independently."""
        sender1 = add_sender('Club Alpha', 'clubalpha@gmail.com', 'pass1')
        sender2 = add_sender('Club Beta', 'clubbeta@gmail.com', 'pass2')
        
        c1 = create_campaign('Alpha Recruitment', sender1['id'])
        c2 = create_campaign('Beta Recruitment', sender2['id'])
        
        self.assertEqual(c1['sender_email'], 'clubalpha@gmail.com')
        self.assertEqual(c2['sender_email'], 'clubbeta@gmail.com')
        self.assertNotEqual(c1['sender_id'], c2['sender_id'])

    def test_03_different_excel_file_isolation(self):
        """TEST 3: Different Excel files remain strictly separated between campaigns."""
        sender = add_sender('Admin', 'admin@gmail.com', 'pass123')
        c1 = create_campaign('Campaign 1', sender['id'])
        c2 = create_campaign('Campaign 2', sender['id'])
        
        c1['excel_filename'] = 'round1.xlsx'
        c1['excel_file_path'] = self.file1
        save_campaign(c1)
        
        c2['excel_filename'] = 'zero.xlsx'
        c2['excel_file_path'] = self.file_zero
        save_campaign(c2)
        
        loaded_c1 = get_campaign_by_id(c1['id'])
        loaded_c2 = get_campaign_by_id(c2['id'])
        
        self.assertEqual(loaded_c1['excel_filename'], 'round1.xlsx')
        self.assertEqual(loaded_c2['excel_filename'], 'zero.xlsx')

    def test_04_missing_email_reported_and_not_sent(self):
        """TEST 4: Missing email rows are marked properly and omitted from dispatch queue."""
        parsed = parse_excel_file(self.file1)
        meera = next(r for r in parsed['rows'] if r['name'] == 'Meera Joshi')
        self.assertEqual(meera['status'], 'Missing Email')
        self.assertFalse(meera['can_send'])

    def test_05_unknown_selection_status_skipped_safely(self):
        """TEST 5: Unknown selection values (Pending, etc.) are skipped and reported."""
        parsed = parse_excel_file(self.file1)
        divya = next(r for r in parsed['rows'] if r['name'] == 'Divya Nair')
        self.assertEqual(divya['status'], 'Skipped — Unknown Selection Status')
        self.assertFalse(divya['can_send'])

    def test_06_invalid_gmail_password_handled_safely(self):
        """TEST 6: Invalid Gmail credentials return user-friendly error without crashing."""
        result = test_gmail_credentials('invalid_user@gmail.com', 'bad_app_password')
        self.assertFalse(result['success'])
        self.assertIn('Unable to authenticate with this Gmail account', result['message'])

    def test_07_one_email_failure_continues_others(self):
        """TEST 7: When sending, if one recipient fails, the loop continues processing others."""
        recipients = [
            {'name': 'Good 1', 'email': 'good1@example.com', 'selection': 'Selected', 'custom_data': {}},
            {'name': 'Bad', 'email': 'bad@example.com', 'selection': 'Selected', 'custom_data': {}},
            {'name': 'Good 2', 'email': 'good2@example.com', 'selection': 'Selected', 'custom_data': {}},
        ]
        
        # Mock SMTP to simulate failure on 2nd recipient only
        with patch('smtplib.SMTP') as mock_smtp_class:
            mock_server = MagicMock()
            mock_smtp_class.return_value = mock_server
            
            # First send succeeds, second raises exception, third succeeds
            mock_server.send_message.side_effect = [None, Exception('Recipient mailbox not found'), None]
            
            sender_info = {
                'email': 'sender@gmail.com',
                'display_name': 'Sender Team',
                'app_password': 'testpassword1234'
            }
            
            res = send_campaign_emails(sender_info, recipients, 'Subject', 'Body')
            
            self.assertEqual(res['sent_count'], 2)
            self.assertEqual(res['failed_count'], 1)
            self.assertEqual(len(res['results']), 3)
            self.assertEqual(res['results'][0]['status'], 'Sent')
            self.assertEqual(res['results'][1]['status'], 'Failed')
            self.assertIn('Recipient mailbox not found', res['results'][1]['error'])
            self.assertEqual(res['results'][2]['status'], 'Sent')

    def test_08_duplicate_send_protection(self):
        """TEST 8: Completed campaign triggers duplicate-send warning before resending."""
        sender = add_sender('Club Head', 'clubhead@gmail.com', 'passwordsupersecret')
        c = create_campaign('Finished Campaign', sender['id'])
        c['status'] = 'completed'
        c['excel_file_path'] = self.file1
        save_campaign(c)
        
        # Sending without force_resend should warn
        res = execute_campaign_send(c['id'], force_resend=False)
        self.assertTrue(res.get('requires_confirmation'))
        self.assertIn('already been processed', res.get('warning'))

    def test_09_no_selected_candidates_aborts_send(self):
        """TEST 9: Campaigns with no selected rows abort safely."""
        sender = add_sender('Judge', 'judge@gmail.com', 'pass')
        c = create_campaign('Zero Campaign', sender['id'])
        c['excel_file_path'] = self.file_zero
        c['mapped_columns'] = {'email_column': 'Email', 'selection_column': 'Selection'}
        c['valid_recipients'] = []
        save_campaign(c)
        
        with self.assertRaises(ValueError) as ctx:
            execute_campaign_send(c['id'], force_resend=False)
        self.assertIn('No valid email addresses were found', str(ctx.exception))

    def test_10_personalization_substitutes_dynamic_variables(self):
        """TEST 10: Dynamic placeholders like {Name}, {Department}, {Round} replace correctly."""
        template_subj = "Hello {Name}, update regarding {Round}"
        template_body = "Dear {Name},\n\nYou have been selected for {Department} in {Round}.\nSent to {Email}."
        
        custom_data = {
            'Name': 'Priya Sharma',
            'Email': 'priya@example.com',
            'Department': 'Cloud Infrastructure',
            'Round': 'Round 1'
        }
        
        sub_subj = substitute_placeholders(template_subj, custom_data)
        sub_body = substitute_placeholders(template_body, custom_data)
        
        self.assertEqual(sub_subj, "Hello Priya Sharma, update regarding Round 1")
        self.assertIn("selected for Cloud Infrastructure", sub_body)
        self.assertIn("Round 1", sub_body)
        self.assertIn("priya@example.com", sub_body)

if __name__ == '__main__':
    unittest.main()
