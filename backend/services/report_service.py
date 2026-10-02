import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'reports')

def ensure_reports_dir():
    os.makedirs(REPORTS_DIR, exist_ok=True)

def generate_campaign_excel_report(campaign: Dict[str, Any], all_rows: List[Dict[str, Any]]) -> str:
    """
    Generates a professional Excel report for a campaign containing:
    - Campaign summary header
    - Table with Name, Email, Selection, Status, Error, Sent Time
    - Professional styling (colors, borders, column auto-widths)
    """
    ensure_reports_dir()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Email Campaign Report"
    
    # Styles
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
    meta_font = Font(name="Segoe UI", size=10, color="64748B")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Segoe UI", size=10, color="1E293B")
    
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid") # Modern Blue
    status_sent_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light Green
    status_failed_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Light Red
    status_skipped_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid") # Light Gray
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    # Title & Metadata
    ws["A1"] = f"Campaign: {campaign.get('name', 'Email Campaign')}"
    ws["A1"].font = title_font
    
    ws["A2"] = f"Sender: {campaign.get('sender_email', 'N/A')} | Excel: {campaign.get('excel_filename', 'N/A')} | Date: {campaign.get('sent_at', campaign.get('created_at', 'N/A'))}"
    ws["A2"].font = meta_font
    
    # Empty line
    start_row = 4
    headers = ["Name", "Email", "Selection Status", "Delivery Status", "Error / Reason", "Sent Time"]
    
    # Write Table Headers
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    # Write Data Rows
    current_row = start_row + 1
    
    # Map results by email or row
    results_map = {}
    for r in campaign.get("results", []):
        results_map[r.get("email", "").lower()] = r
        
    for item in all_rows:
        email_clean = item.get("email", "").lower()
        res_info = results_map.get(email_clean)
        
        # Determine status and error
        if res_info:
            delivery_status = res_info.get("status", "Unknown")
            error_msg = res_info.get("error", "")
            sent_time = res_info.get("sent_at", "")
        else:
            delivery_status = item.get("status", "Skipped")
            error_msg = "Not included in sending queue" if item.get("selection") != "Selected" else item.get("status", "")
            sent_time = ""
            
        row_values = [
            item.get("name", ""),
            item.get("email", ""),
            item.get("raw_selection", item.get("selection", "")),
            delivery_status,
            error_msg,
            sent_time
        ]
        
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            
            # Badge fill for Delivery Status column
            if col_idx == 4:
                if delivery_status == "Sent":
                    cell.fill = status_sent_fill
                elif delivery_status == "Failed":
                    cell.fill = status_failed_fill
                else:
                    cell.fill = status_skipped_fill
                    
        current_row += 1
        
    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            if cell.row >= start_row and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
        
    # Generate file path
    safe_name = "".join([c if c.isalnum() else "_" for c in campaign.get("name", "campaign")])
    report_filename = f"{safe_name}_Report_{campaign.get('id', 'export')[:8]}.xlsx"
    report_path = os.path.join(REPORTS_DIR, report_filename)
    
    wb.save(report_path)
    return report_path
